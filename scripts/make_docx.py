"""Convert paper/main.tex to a Word file (paper/main.docx).

Pipeline: expand the number macros (paper/numbers.tex), inline the generated tables, number tables/figures and resolve \\ref,
convert with pandoc (citations via citeproc + paper/refs.bib), then tidy styles and tables with python-docx.

  python scripts/make_docx.py            # writes paper/main.docx
"""
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
P = ROOT / "paper"
WORK = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/tmp/make_docx")
WORK.mkdir(parents=True, exist_ok=True)

CSL = """<?xml version="1.0" encoding="utf-8"?>
<style xmlns="http://purl.org/net/xbiblio/csl" class="in-text" version="1.0" default-locale="en-US">
  <info><title>Compact numeric</title><id>http://example.org/compact-numeric</id><updated>2026-01-01T00:00:00+00:00</updated></info>
  <macro name="author"><names variable="author"><name and="text" initialize-with="." delimiter=", " delimiter-precedes-last="never" name-as-sort-order="all" sort-separator=" "/></names></macro>
  <citation collapse="citation-number">
    <sort><key variable="citation-number"/></sort>
    <layout prefix="[" suffix="]" delimiter=", "><text variable="citation-number"/></layout>
  </citation>
  <bibliography entry-spacing="0" second-field-align="flush">
    <layout>
      <text variable="citation-number" prefix="[" suffix="]"/>
      <text macro="author" prefix=" " suffix=". "/>
      <text variable="title" suffix=". "/>
      <text variable="container-title" font-style="italic" suffix=", "/>
      <text variable="publisher" suffix=", "/>
      <group delimiter=", " suffix=", ">
        <text variable="volume"/><text variable="page"/>
      </group>
      <date variable="issued"><date-part name="year"/></date>
      <text variable="note" prefix=". "/>
      <text variable="medium" prefix=". "/>
      <text variable="URL" prefix=". "/>
    </layout>
  </bibliography>
</style>
"""


MATH_TEXT = {
    "K": r"\emph{K}", "z": r"\emph{z}", "s": r"\emph{s}", "k": r"\emph{k}", "y": r"\emph{y}", "x": r"\emph{x}", "q": r"\emph{q}",
    r"R_i": r"\emph{R}\textsubscript{\emph{i}}",
    r"P_k(X\mid Y)": r"\emph{P}\textsubscript{\emph{k}}(\emph{X} | \emph{Y})",
    r"P_k(Y)": r"\emph{P}\textsubscript{\emph{k}}(\emph{Y})",
    r"P(\text{COVID-19})": r"\emph{P}(COVID-19)",
    r"\mathcal{L}_{\text{adv}}=\mathbb{E}\,\mathrm{CE}(d_y(\mathrm{GRL}_\lambda(z)),s)":
        r"\emph{L}\textsubscript{adv} = E CE(\emph{d}\textsubscript{\emph{y}}(GRL\textsubscript{λ}(\emph{z})), \emph{s})",
    r"w_{s,y}=\pi(y)/\pi_s(y)": r"\emph{w}\textsubscript{\emph{s,y}} = π(\emph{y}) / π\textsubscript{\emph{s}}(\emph{y})",
    r"1-p_\theta(y\mid x)": r"1 − \emph{p}\textsubscript{θ}(\emph{y} | \emph{x})",
    r"\{y: 1-p_\theta(y\mid x)\le\hat q\}": "\\{\\emph{y}: 1 − \\emph{p}\\textsubscript{θ}(\\emph{y} | \\emph{x}) ≤ \\emph{q}\u0302\\}",
    r"\#\{i: R_i<1-p_{-k(i)}(y\mid x)\}<(1-\alpha)(n+1)":
        r"\#\{\emph{i}: \emph{R}\textsubscript{\emph{i}} < 1 − \emph{p}\textsubscript{−\emph{k}(\emph{i})}(\emph{y} | \emph{x})\} < (1 − α)(\emph{n} + 1)",
    r"\alpha=0.1": "α = 0.1",
    r"|\text{coverage}-(1-\alpha)|": "|coverage − (1 − α)|",
    r"\alpha": "α", r"\Delta": "Δ",
    r"64\times64": "64×64",
    r"2\times10^{-3}": r"2×10\textsuperscript{−3}",
    r"0.662\to0.660": "0.662 → 0.660",
    "V0,V2,V4": "V0, V2, V4", "<": "<", "<0.85": "<0.85",
    "^{d}": r"\textsuperscript{d}", "^{c}": r"\textsuperscript{c}",
}


def simplify_math(tex: str) -> str:
    """Turn inline math into plain Unicode text so it renders everywhere (no Word equation objects)."""
    def rep(m):
        body = m.group(1)
        if re.fullmatch(r"[\d\.\s\+\-\u2212,\[\]\(\)%/\\]*", body):  # numbers, signs, brackets, \%
            return body.replace("-", "\u2212")
        if body in MATH_TEXT:
            return MATH_TEXT[body]
        if body.endswith("\\%") and re.fullmatch(r"[\d\.]+\\%", body):
            return body
        print("WARNING: math left as equation:", body)
        return m.group(0)
    return re.sub(r"(?<!\\)\$(.+?)(?<!\\)\$", rep, tex, flags=re.S)


def expand_tex() -> tuple[str, str, str]:
    tex = (P / "main.tex").read_text()
    nums = {m.group(1): m.group(2) for m in re.finditer(r"\\newcommand\{\\(\w+)\}\{(.*?)\}\n", (P / "numbers.tex").read_text())}
    nums = {k: v.replace("\\ensuremath{-}", "\u2212") for k, v in nums.items()}
    # title block -> metadata
    m = re.search(r"\\title\{(.*?)\}\n\\author\{(.*?)\}\n\\date\{\}", tex, re.S)
    title_raw, author = m.group(1), m.group(2)
    main_title = re.sub(r"\\vspace\{[^}]*\}", "", title_raw.split("\\\\")[0]).strip()
    subtitle = re.sub(r"\\large\s*", "", title_raw.split("\\\\", 1)[1]).strip()
    subtitle = re.sub(r"^\[[^\]]*\]", "", subtitle).strip()
    tex = tex[:m.start()] + tex[m.end():]
    tex = tex.replace("\\maketitle", "").replace("\\input{numbers}", "")
    tex = re.sub(r"\\vspace\{[^}]*\}", "", tex)
    tex = tex.replace("\\setlength\\itemsep{0pt}", "").replace("\\allowbreak ", "").replace("\\allowbreak", "")
    # drop bibliography block; pandoc appends it. Move appendix before it.
    tex = re.sub(r"\{\\small\s*\\bibliographystyle\{plainnat\}\s*\\bibliography\{refs\}\}", "", tex)
    # remove resizebox wrappers, inline tables
    tex = re.sub(r"\\resizebox\{\\linewidth\}\{!\}\{(\\input\{[^}]*\})\}", r"\1", tex)
    tex = re.sub(r"\\input\{(tables/[^}]*)\}", lambda mm: (P / mm.group(1)).read_text(), tex)
    # figures: png instead of pdf
    tex = tex.replace(".pdf}", ".png}")
    # number floats in order of appearance and resolve refs
    labels, tcount, fcount = {}, 0, 0
    out, pos = [], 0
    for fm in re.finditer(r"\\begin\{(table|figure)\}(.*?)\\end\{\1\}", tex, re.S):
        kind, body = fm.group(1), fm.group(2)
        if kind == "table":
            tcount += 1; n, name = tcount, "Table"
        else:
            fcount += 1; n, name = fcount, "Figure"
        lab = re.search(r"\\label\{([^}]*)\}", body)
        if lab:
            labels[lab.group(1)] = (name, n)
        body = re.sub(r"\\label\{[^}]*\}", "", body)
        body = body.replace("\\caption{", "\\caption{\\textbf{%s %d.} " % (name, n), 1)
        out.append(tex[pos:fm.start()]); out.append("\\begin{%s}%s\\end{%s}" % (kind, body, kind)); pos = fm.end()
    out.append(tex[pos:]); tex = "".join(out)
    tex = re.sub(r"(Table|Fig\.|Figure)~\\ref\{([^}]*)\}", lambda mm: "%s %d" % (labels[mm.group(2)][0], labels[mm.group(2)][1]), tex)
    tex = re.sub(r"\\ref\{([^}]*)\}", lambda mm: str(labels[mm.group(1)][1]), tex)
    tex = tex.replace("\\captionsetup", "%\\captionsetup")
    # expand number macros (longest first, not followed by a letter)
    for k in sorted(nums, key=len, reverse=True):
        tex = re.sub(r"\\%s(?![A-Za-z])" % k, lambda _m, v=nums[k]: v, tex)
    tex = simplify_math(tex)
    left = re.findall(r"\\[A-Za-z]+Cov[A-Za-z]*|\\T[A-Z][A-Za-z]+|\\rel[A-Z][A-Za-z]+", tex)
    assert not left, left[:5]
    return tex, main_title, subtitle


def make_reference_doc(path: Path):
    ref = WORK / "ref_default.docx"
    subprocess.run(["pandoc", "-o", str(ref), "--print-default-data-file", "reference.docx"], check=True)
    d = Document(str(ref))
    for s in d.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Inches(0.95); s.top_margin = s.bottom_margin = Inches(0.9)
    st = d.styles
    def font(name, size=None, bold=None, italic=None, color=None, align=None, before=None, after=None):
        s = next((x for x in st if x.name and x.name.lower() == name.lower()), None)
        if s is None:
            print("style not found:", name)
            return
        s.font.name = "Times New Roman"
        rpr = s.element.get_or_add_rPr()
        rf = rpr.rFonts
        if rf is not None:
            for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                rf.attrib.pop(qn("w:" + a), None)
            for a in ("eastAsia", "cs", "hAnsi", "ascii"):
                rf.set(qn("w:" + a), "Times New Roman")
        if size: s.font.size = Pt(size)
        if bold is not None: s.font.bold = bold
        if italic is not None: s.font.italic = italic
        s.font.color.rgb = RGBColor(0, 0, 0)
        col = rpr.find(qn("w:color"))
        if col is not None:
            for a in ("themeColor", "themeShade", "themeTint"):
                col.attrib.pop(qn("w:" + a), None)
        if align is not None: s.paragraph_format.alignment = align
        if before is not None: s.paragraph_format.space_before = Pt(before)
        if after is not None: s.paragraph_format.space_after = Pt(after)
    for n in ("Normal", "Body Text", "First Paragraph", "Compact"):
        font(n, 10.5, before=0, after=5)
        sty = next((x for x in st if x.name and x.name.lower() == n.lower()), None)
        if sty is not None:
            sty.paragraph_format.line_spacing = 1.0
    font("Title", 18, True, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6)
    font("Subtitle", 13, False, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    font("Author", 11, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    font("Date", 11, align=WD_ALIGN_PARAGRAPH.CENTER)
    font("Abstract", 10, after=6)
    font("Abstract Title", 11, True, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=3)
    font("Heading 1", 13, True, before=10, after=4)
    font("Heading 2", 11.5, True, before=8, after=3)
    font("Heading 3", 11, True, True, before=8, after=3)
    font("Image Caption", 9.5, False, False, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    font("Table Caption", 9.5, False, False, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=4)
    font("Captioned Figure", 11, align=WD_ALIGN_PARAGRAPH.CENTER)
    font("Bibliography", 9, after=1)
    d.save(str(path))


def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge, sz in kw.items():
        el = OxmlElement("w:%s" % edge)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        borders.append(el)


def set_table_layout(t, widths_dxa):
    tbl = t._tbl
    tblPr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(sum(widths_dxa))); w.set(qn("w:type"), "dxa"); tblPr.append(w)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for wd in widths_dxa:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(wd)); grid.append(gc)
    for row in t.rows:
        for cell, wd in zip(row.cells, widths_dxa):
            tcPr = cell._tc.get_or_add_tcPr()
            for el in tcPr.findall(qn("w:tcW")):
                tcPr.remove(el)
            tw = OxmlElement("w:tcW"); tw.set(qn("w:w"), str(wd)); tw.set(qn("w:type"), "dxa"); tcPr.insert(0, tw)


def tidy(path: Path):
    d = Document(str(path))
    sec = d.sections[0]
    text_w = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)  # EMU -> DXA
    sec.header_distance = Inches(0.45); sec.footer_distance = Inches(0.45); sec.gutter = 0
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for kind, txt in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        run._r.append(el)
    run.font.size = Pt(9); run.font.name = "Times New Roman"
    # pandoc wraps each figure in a layout table; replace it by a plain centred paragraph holding the picture
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for tbl in list(d.element.body.iter("{%s}tbl" % W)):
        if tbl.find(".//{%s}drawing" % W) is None:
            continue
        para = tbl.find(".//{%s}drawing/../.." % W)  # w:p containing the drawing run
        tbl.addprevious(para)
        tbl.getparent().remove(tbl)
    for t in d.tables:
        for tn in t._tbl.iter(qn("w:t")):
            if tn.text and "\u00a0" in tn.text:
                tn.text = tn.text.replace("\u00a0", " ")
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        ncols = len(t.columns)
        size = Pt(8) if ncols >= 7 else Pt(9)
        # column widths: proportional to the longest unbreakable chunk / text in each column
        want = []
        for ci in range(ncols):
            longest = 0
            for ri, row in enumerate(t.rows):
                txt = row.cells[ci].text
                words = txt.split() or [""]
                if ci == 0:
                    longest = max(longest, min(len(txt), 30))
                else:
                    longest = max(longest, max(len(w) for w in words) + 3, min(len(txt), 22) if ri > 0 else 0)
            want.append(max(longest, 7))
        tot = sum(want)
        widths = [int(text_w * w / tot) for w in want]
        widths[-1] += text_w - sum(widths)
        set_table_layout(t, widths)
        mar = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            m_ = OxmlElement("w:" + side); m_.set(qn("w:w"), "50"); m_.set(qn("w:type"), "dxa"); mar.append(m_)
        t._tbl.tblPr.append(mar)
        nrows = len(t.rows)
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    if ri < nrows - 1 and nrows <= 14:
                        p.paragraph_format.keep_with_next = True
                    for r in p.runs:
                        r.font.size = size; r.font.name = "Times New Roman"
                        for tnode in r._r.findall(qn("w:t")):
                            tnode.text = re.sub(r"(?<![\w\u2212])-(?=\d)", "\u2212", tnode.text or "")
                        if ri == 0: r.font.bold = True
                if ri == 0:
                    set_cell_border(cell, top=8, bottom=6)
                if ri == len(t.rows) - 1:
                    set_cell_border(cell, bottom=8)
    order = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
             "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]
    for t in d.tables:
        pr = t._tbl.tblPr
        kids = sorted(list(pr), key=lambda e: order.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in order else 99)
        for k in kids:
            pr.remove(k)
        for k in kids:
            pr.append(k)
    for t in d.tables:
        nxt = t._tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            Paragraph(nxt, d._body).paragraph_format.space_before = Pt(8)
    # figures: set width as a fraction of the text width (aspect ratio kept) and centre them
    from docx.shared import Emu
    frac = {"fig_diagnostics": 0.88, "fig_methods": 0.80, "fig_conformal": 0.98, "fig_optimism": 0.62}
    text_emu = sec.page_width - sec.left_margin - sec.right_margin
    ns = {"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main", "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture"}
    for para in d.paragraphs:
        for inl in para._p.iter("{%s}inline" % ns["wp"]):
            cnv = inl.find(".//pic:cNvPr", ns)
            name = cnv.get("descr") if cnv is not None else ""
            blip = inl.find(".//a:blip", ns)
            rid = blip.get(qn("r:embed")) if blip is not None else None
            target = d.part.related_parts[rid].partname if rid else ""
            key = next((k for k in frac if k in (name or "") or k in str(target)), None)
            ext = inl.find("wp:extent", ns)
            if key is None:
                print("figure not matched:", name, target); continue
            cx, cy = int(ext.get("cx")), int(ext.get("cy"))
            new_cx = int(text_emu * frac[key]); new_cy = int(cy * new_cx / cx)
            ext.set("cx", str(new_cx)); ext.set("cy", str(new_cy))
            for e in inl.findall(".//pic:spPr/a:xfrm/a:ext", ns):
                e.set("cx", str(new_cx)); e.set("cy", str(new_cy))
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.keep_with_next = True
    d.save(str(path))


def main():
    tex, title, subtitle = expand_tex()
    (WORK / "paper_expanded.tex").write_text(tex)
    (WORK / "numeric.csl").write_text(CSL)
    ref = WORK / "reference.docx"; make_reference_doc(ref)
    out = P / "main.docx"
    cmd = ["pandoc", str(WORK / "paper_expanded.tex"), "-f", "latex+raw_tex", "-t", "docx", "--citeproc", "--bibliography", str(P / "refs.bib"),
           "--csl", str(WORK / "numeric.csl"), "--reference-doc", str(ref), "--resource-path", str(P), "--number-sections",
           "-M", f"title={title}", "-M", f"subtitle={subtitle}", "-M", "author=Anonymous Authors",
           "-M", "reference-section-title=References", "-o", str(out)]
    r = subprocess.run(cmd, capture_output=True, text=True)
    print(r.stderr[-1500:])
    r.check_returncode()
    tidy(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
